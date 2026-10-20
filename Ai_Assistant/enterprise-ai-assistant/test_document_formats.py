import asyncio
import io
import json
import csv

from fastapi import UploadFile
from openpyxl import Workbook
from docx import Document

from app.config.settings import get_settings
from app.services.document_service import DocumentService


async def test_file(service, filename, content):
    upload = UploadFile(
        filename=filename,
        file=io.BytesIO(content),
    )

    result = await service.read_document(upload)

    print(f"\n--- {filename} ---")
    print("SUCCESS")
    print("File type:", result.file_type)
    print("Text:", result.text[:200])
    print("Words:", result.word_count)


async def main():
    service = DocumentService(get_settings())

    # TXT
    await test_file(
        service,
        "test.txt",
        b"Enterprise AI Assistant TXT document test.",
    )

    # MD
    await test_file(
        service,
        "test.md",
        b"# Enterprise AI Assistant\n\nMarkdown document test.",
    )

    # JSON
    json_content = json.dumps({
        "project": "Enterprise AI Assistant",
        "status": "working",
        "type": "AI Assistant",
    }).encode()

    await test_file(
        service,
        "test.json",
        json_content,
    )

    # CSV
    csv_buffer = io.StringIO()
    writer = csv.writer(csv_buffer)
    writer.writerow(["Name", "Role"])
    writer.writerow(["Shruti", "Developer"])
    writer.writerow(["AI Assistant", "Project"])

    await test_file(
        service,
        "test.csv",
        csv_buffer.getvalue().encode(),
    )

    # DOCX
    doc = Document()
    doc.add_paragraph("Enterprise AI Assistant DOCX test.")
    doc.add_paragraph("Document processing is working.")
    
    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)

    await test_file(
        service,
        "test.docx",
        doc_buffer.getvalue(),
    )

    # XLSX
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Project"

    sheet.append(["Project", "Status"])
    sheet.append(["Enterprise AI Assistant", "Working"])

    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)

    await test_file(
        service,
        "test.xlsx",
        xlsx_buffer.getvalue(),
    )


if __name__ == "__main__":
    asyncio.run(main())